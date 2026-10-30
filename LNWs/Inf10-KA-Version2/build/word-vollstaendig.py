"""Rebuild the retained PDF as four Word pages with editable text.

The original vector furniture is rasterized at 300 dpi without any text.
All PDF text is recreated as positioned Word text boxes, in line-sized groups.
"""
from pathlib import Path
from collections import defaultdict
from hashlib import sha256
import json, math, subprocess, unicodedata
from lxml import etree
from PIL import ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ContentStream, NameObject
import pdfplumber
from docx import Document
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
build = root / 'build/word-vollstaendig'
build.mkdir(exist_ok=True)
source = root / 'Inf10-KA-Version2-loesungen.pdf'
output = root / 'Inf10-KA-Version2-loesungen-vollstaendig.docx'
reference_hash = sha256(source.read_bytes()).hexdigest()
poppler = Path('C:/Users/Franke/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin')
reader = PdfReader(source)
writer = PdfWriter()
for page in reader.pages:
    writer.add_page(page)
for page in writer.pages:
    content = ContentStream(page.get_contents(), writer)
    content.operations = [(args, op) for args, op in content.operations if op not in (b'Tj', b'TJ', b"'", b'"')]
    page[NameObject('/Contents')] = writer._add_object(content)
background_pdf = build / 'grafiken-ohne-text.pdf'
with background_pdf.open('wb') as stream:
    writer.write(stream)
assert all(not p.extract_text().strip() for p in PdfReader(background_pdf).pages)
subprocess.run([str(poppler / 'pdftoppm.exe'), '-r', '300', '-png', str(background_pdf), str(build / 'hintergrund')], check=True)

NS_V = 'urn:schemas-microsoft-com:vml'
NS_O = 'urn:schemas-microsoft-com:office:office'
NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS_W10 = 'urn:schemas-microsoft-com:office:word'
def v(tag): return '{' + NS_V + '}' + tag
def node(tag, attrs=None):
    el = OxmlElement(tag)
    for k, val in (attrs or {}).items(): el.set(qn(k), str(val))
    return el

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(0)
sec.header_distance = sec.footer_distance = Mm(0)
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(1)
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.core_properties.title = 'Musterlösung Informatik Datenbanken Version 2 mit vollständiger Angabe'
doc.core_properties.author = 'Franke'
doc.core_properties.subject = 'Kurzarbeit Klasse 10 vom 27.10.2026'

fonts = {}
def font_properties(pdf_font):
    bold = 'Bold' in pdf_font
    italic = any(s in pdf_font for s in ('Oblique', 'Italic'))
    mono = any(s in pdf_font for s in ('Mono', 'Typewriter'))
    family = 'Courier New' if mono else 'Arial'
    if mono:
        filename = 'courbi.ttf' if bold and italic else 'courbd.ttf' if bold else 'couri.ttf' if italic else 'cour.ttf'
    else:
        filename = 'arialbi.ttf' if bold and italic else 'arialbd.ttf' if bold else 'ariali.ttf' if italic else 'arial.ttf'
    return family, bold, italic, filename

def word_width(text, filename, size):
    key = (filename, size)
    if key not in fonts:
        fonts[key] = ImageFont.truetype('C:/Windows/Fonts/' + filename, round(size * 100))
    return fonts[key].getlength(text) / 100

def background(paragraph, path, width, height, page_num):
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Pt(width), height=Pt(height))._inline
    anchor = etree.Element(qn('wp:anchor'), nsmap=inline.nsmap)
    for k, val in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'1','layoutInCell':'0','allowOverlap':'1'}.items():
        anchor.set(k, val)
    etree.SubElement(anchor, qn('wp:simplePos'), x='0', y='0')
    for axis in ('H','V'):
        pos = etree.SubElement(anchor, qn('wp:position' + axis), relativeFrom='page')
        etree.SubElement(pos, qn('wp:posOffset')).text = '0'
    for tag in ('wp:extent', 'wp:effectExtent'):
        child = inline.find(qn(tag))
        if child is not None: anchor.append(child)
    etree.SubElement(anchor, qn('wp:wrapNone'))
    for tag in ('wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
        child = inline.find(qn(tag))
        if child is not None:
            if tag == 'wp:docPr':
                child.set('name', f'Grafiken und Raster Seite {page_num}')
                child.set('descr', 'Originale Tabellenlinien, zweigeteilte Klassenkarten, Verbindungen und Antwortbereiche. Sämtliche Beschriftungen stehen als bearbeitbarer Word-Text darüber.')
            anchor.append(child)
    inline.getparent().replace(inline, anchor)

def add_run(p, text, word, width=None):
    if not text: return
    family, bold, italic, filename = font_properties(word['fontname'])
    size = round(word['size'] * 2) / 2
    r = node('w:r')
    props = node('w:rPr')
    props.append(node('w:rFonts', {'w:ascii':family, 'w:hAnsi':family, 'w:eastAsia':family, 'w:cs':family}))
    props.append(node('w:sz', {'w:val':round(size * 2)}))
    props.append(node('w:szCs', {'w:val':round(size * 2)}))
    props.append(node('w:lang', {'w:val':'de-DE'}))
    if bold: props.append(node('w:b'))
    if italic: props.append(node('w:i'))
    color = word.get('non_stroking_color', (0,0,0))
    if not isinstance(color, (tuple, list)): color = (color, color, color)
    if len(color) < 3: color = (color[0],) * 3
    rgb = ''.join(f'{max(0,min(255,round(c * 255))):02X}' for c in color[:3])
    props.append(node('w:color', {'w:val':rgb}))
    if width is not None:
        natural = word_width(text, filename, size)
        if natural:
            scale = max(1, min(600, round(width / natural * 100)))
            props.append(node('w:w', {'w:val':scale}))
    r.append(props)
    t = node('w:t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    r.append(t)
    p.append(r)

all_groups = []
shape_id = 1
with pdfplumber.open(source) as pdf:
    for page_num, page in enumerate(pdf.pages, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(1)
        if page_num > 1: p.paragraph_format.page_break_before = True
        background(p, build / f'hintergrund-{page_num}.png', float(page.width), float(page.height), page_num)
        words = page.extract_words(extra_attrs=['fontname', 'size', 'non_stroking_color'], return_chars=True, x_tolerance=1.3, y_tolerance=1.2, split_at_punctuation=False)
        lines = defaultdict(list)
        for word in words:
            baseline = float(page.height) - word['chars'][0]['matrix'][5]
            key = round(baseline * 2) / 2
            word['baseline'] = baseline
            lines[key].append(word)
        for key in sorted(lines):
            groups = []
            for word in sorted(lines[key], key=lambda w:w['x0']):
                if groups and word['x0'] - groups[-1][-1]['x1'] <= 13:
                    groups[-1].append(word)
                else:
                    groups.append([word])
            for group in groups:
                x = group[0]['x0']
                baseline = group[0]['baseline']
                size = max(w['size'] for w in group)
                y = baseline - size * 1.05
                width = group[-1]['x1'] - x + 4
                height = size * 2.2
                r = p.add_run()._r
                pict = node('w:pict')
                shape = etree.Element(v('shape'), nsmap={'v':NS_V,'o':NS_O,'w10':NS_W10})
                shape.set('id', f'_x0000_s{shape_id}')
                shape.set('type', '#_x0000_t202')
                shape.set('filled','f'); shape.set('stroked','f')
                shape.set('style', f'position:absolute;left:{x:.4f}pt;top:{y:.4f}pt;width:{width:.4f}pt;height:{height:.4f}pt;z-index:{shape_id+10};mso-position-horizontal-relative:page;mso-position-vertical-relative:page;mso-wrap-style:none')
                textbox = etree.SubElement(shape, v('textbox'), inset='0,0,0,0', style='mso-fit-shape-to-text:f')
                content = node('w:txbxContent'); textbox.append(content)
                paragraph = node('w:p'); content.append(paragraph)
                pr = node('w:pPr')
                pr.append(node('w:spacing', {'w:before':0,'w:after':0,'w:line':round(size*1.2*20),'w:lineRule':'exact'}))
                pr.append(node('w:widowControl', {'w:val':0}))
                paragraph.append(pr)
                for i, word in enumerate(group):
                    if i:
                        gap = word['x0'] - group[i-1]['x1']
                        if gap > .6:
                            add_run(paragraph, ' ', word, gap)
                    text = unicodedata.normalize('NFKC', word['text'])
                    add_run(paragraph, text, word, word['x1'] - word['x0'])
                pict.append(shape); r.append(pict)
                all_groups.append({'page':page_num,'id':shape_id,'x':x,'y':y,'baseline':baseline,'text':' '.join(w['text'] for w in group)})
                shape_id += 1

doc.save(output)
assert sha256(source.read_bytes()).hexdigest() == reference_hash
(build / 'textfelder.json').write_text(json.dumps(all_groups, ensure_ascii=False, indent=2), encoding='utf-8')
(build / 'artifact.md').write_text(f'''# Ausführungskontrakt
Referenz: {source.resolve()}
SHA256: {reference_hash}
Vier A4-Seiten. Vollständiger Aufgaben- und Lösungstext aus der PDF bleibt erhalten.
Seitentreue Rekonstruktion mit editierbaren Word-Textfeldern je Zeile bzw. zusammenhängendem Teil einer Zeile. Native Zeichenformatierung übernimmt Größe, Fettdruck, Kursivschrift und Farbe. Latin Modern Sans wird durch Arial, Latin Modern Mono durch Courier New angenähert. Die Laufweite wird der Referenz angepasst.
PDF-Grafiken ohne Text als Hintergrundgrafik mit 300 dpi: Linien, Tabellenraster, Karoflächen, Klassenkarten, Verbindungen und Durchstreichungen. Keine vollständigen Seitenbilder: der gesamte Text bleibt editierbar.
Deckblatt: Titel22pt, Kopfdaten14pt, Datum27.10.2026. Bewertungsverteilung4/11/8/6, Gesamt29BE. Seite2 beginnt mit Aufgabe1, Seite3 mit Aufgabe3, Seite4 mit3c. Seitenzahlen und Schlussgruß übernehmen.
Angabentext, fehlerhafte SQL-Vorgabe, ursprüngliches und korrigiertes Diagramm aus der PDF unverändert. Die einfache Word-Fassung und alle PDF/LaTeX-Dateien bleiben unverändert.
Abnahme: Textabgleich über OOXML und gerendertes PDF, vier Seiten, Sichtprüfung jeder Seite.
''',encoding='utf-8')
print(json.dumps({'output':str(output.resolve()),'pages':len(reader.pages),'textboxes':shape_id-1}))
