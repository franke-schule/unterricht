from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root = Path(__file__).resolve().parent.parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(18)
sec.left_margin = sec.right_margin = Mm(20)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(5)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(19)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.core_properties.title = 'Musterlösung Informatik Datenbanken Version 2'
doc.core_properties.author = 'Franke'
doc.core_properties.subject = 'Kurzarbeit Klasse 10 am 27.10.2026'
lang = OxmlElement('w:lang')
lang.set(qn('w:val'), 'de-DE')
doc.styles['Normal'].element.get_or_add_rPr().append(lang)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)

def para(text, bold=False):
    p = doc.add_paragraph()
    p.add_run(text).bold = bold
    return p

def heading(text, level=1):
    return doc.add_heading(text, level)

def code(text, crossed=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
    r.font.strike = crossed
    return p

def table(headers, records, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Mm(width)
        cell.text = text
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'EAF3FA')
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for record in records:
        for cell, text, width in zip(t.add_row().cells, record, widths):
            cell.width = Mm(width)
            cell.text = str(text)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                edge_el = OxmlElement('w:' + edge)
                for key, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                    edge_el.set(qn('w:' + key), value)
                borders.append(edge_el)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '85')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc.add_paragraph('Musterlösung Informatik Datenbanken', 'Title')
para('Kurzarbeit Klasse 10 · Version 2 · 27.10.2026 · Gesamt 29 BE')

heading('1 Grundwissen')
para('4 BE', True)
for text in (
    'Eine Spalte entspricht in der objektorientierten Programmierung einem Attribut.',
    'Ein Datensatz entspricht in der objektorientierten Programmierung einem Objekt.',
    'Eine Tabelle entspricht in der objektorientierten Programmierung einer Klasse.',
    'Eine Zelle - außer in der ersten Zeile - entspricht in der objektorientierten Programmierung einem Attributwert.',
):
    para(text)

heading('2 Apps auf dem Handy')
para('11 BE', True)
heading('2a Klassendiagramm Nr 1', 2)
table(['Klasse', 'Kardinalität', 'Beziehung', 'Kardinalität', 'Klasse'], [
    ('App', 'n', 'ist installiert auf', 'm', 'Handy'),
    ('Handy', 'n', 'gehört zu', '1', 'Benutzer'),
], [29, 24, 64, 24, 29])

heading('2b Beziehungen benennen und begründen', 2)
para('App–Handy: n:m-Beziehung. Eine App kann auf mehreren Handys installiert sein; ein Handy kann mehrere Apps enthalten.')
para('Benutzer–Handy: 1:n-Beziehung. Eine Person kann mehrere Handys haben; jedes Handy gehört genau einer Person. Daher steht im Diagramm n bei Handy und 1 bei Benutzer.')

heading('2c Beziehungstabelle', 2)
para('Eine Beziehungstabelle ist bei einer n:m-Beziehung notwendig. Ein einzelner Fremdschlüssel in einer der beiden Tabellen reicht für die vielen Zuordnungen in beiden Richtungen nicht aus. Jede Zeile der Beziehungstabelle verbindet über zwei Fremdschlüssel genau zwei Datensätze.')

doc.add_page_break()
heading('3 Berechtigungen und Apps')
para('8 BE', True)
heading('3a Neue Berechtigung für WhatsApp', 2)
para('WhatsApp hat die id 1; Standort hat die id 2. In Berechtigungen_von_Apps wird eine neue Zeile mit app_id = 1 und berechtigung_id = 2 ergänzt.')
table(['app_id', 'berechtigung_id'], [(1, 2)], [40, 55])
para('Beide Stammdatensätze existieren bereits; in Apps und Berechtigung sind keine neuen Zeilen erforderlich.')

heading('3b Korrigiertes Klassendiagramm Nr 2', 2)
table(['Klasse', 'Kardinalität', 'Beziehung', 'Kardinalität', 'Klasse'], [
    ('Berechtigungen', 'n', 'ist freigegeben für', 'm', 'Apps'),
], [35, 24, 64, 24, 23])
para('Die korrigierte Skizze enthält die beiden Klassen Berechtigungen und Apps mit einer direkten n:m-Beziehung. Die zusätzliche Klassenkarte Berechtigungen-von-Apps entfällt.')

heading('3c SQL korrigieren', 2)
para('SELECT * bleibt erhalten. Die FROM- und WHERE-Zeilen der Vorgabe werden durchgestrichen und wie folgt korrigiert:')
code('SELECT *\nFROM Berechtigung, Apps, Berechtigungen_von_Apps\nWHERE Berechtigung.id = Berechtigungen_von_Apps.berechtigung_id\n  AND Apps.id = Berechtigungen_von_Apps.app_id;')

heading('4 Appentwickler')
para('6 BE', True)
heading('4a SQL beschreiben', 2)
code('SELECT COUNT(*)\nFROM App\nWHERE erscheinungsjahr > 2025')
para('Die Anweisung zählt die Apps, die nach dem Jahr 2025 erschienen sind, also ab 2026. Ausgegeben wird eine einzelne Zahl. Apps mit Erscheinungsjahr 2025 werden nicht mitgezählt.')

heading('4b Namen von Apps und Appentwicklern ausgeben', 2)
code('SELECT App.name, Appentwickler.name\nFROM App, Appentwickler\nWHERE App.appentwickler_id = Appentwickler.id;')
para('Die beiden Tabellen werden über den Fremdschlüssel App.appentwickler_id und den Primärschlüssel Appentwickler.id verknüpft.')
para('Gesamt 29 BE', True)
para('Viel Erfolg! Franke')

out = root / 'Inf10-KA-Version2-loesungen.docx'
doc.save(out)
check = Document(out)
text = '\n'.join(p.text for p in check.paragraphs)
for required in ('27.10.2026', 'Ein Datensatz entspricht', 'SELECT COUNT(*)', 'WHERE App.appentwickler_id = Appentwickler.id;', 'Berechtigungen_von_Apps.berechtigung_id', 'Gesamt 29 BE'):
    assert required in text
assert 'JOIN' not in text
assert len(check.tables) == 3
print(out.resolve())
