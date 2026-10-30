"""Full statement and solutions in normal Word paragraphs, without grids."""
from pathlib import Path
import subprocess, re, json
from PIL import Image
import pdfplumber
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parent.parent
build = root / 'build/word-schlicht'
build.mkdir(exist_ok=True)

# Reuse the plain Word styles and native table helpers from the first export.
helper = root / 'build/word-erstellen.py'
setup = helper.read_text(encoding='utf-8').split("doc.add_paragraph('Musterlösung Informatik Datenbanken', 'Title')", 1)[0]
exec(compile(setup, str(helper), 'exec'), globals())
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.core_properties.title = 'Musterlösung Informatik Datenbanken mit vollständiger Angabe'

poppler = Path('C:/Users/Franke/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe')
subprocess.run([str(poppler), '-r', '220', '-png', str(root/'Inf10-KA-Version2-loesungen.pdf'), str(build/'vorlage')], check=True)
with pdfplumber.open(root/'Inf10-KA-Version2-loesungen.pdf') as pdf:
    p2 = pdf.pages[1]
    p3 = pdf.pages[2]
    bold2 = p2.extract_words(extra_attrs=['fontname'])
    bold3 = p3.extract_words(extra_attrs=['fontname'])
    app = next(w for w in bold2 if w['text']=='App' and 'Bold' in w['fontname'])
    middle = next(w for w in bold3 if w['text']=='Berechtigungen-von-Apps' and 'Bold' in w['fontname'])
    corrected = [w for w in bold3 if w['text']=='Berechtigungen' and 'Bold' in w['fontname']][-1]
    for name, page, word in (('diagramm1',2,app),('diagramm2',3,middle),('diagramm2-korrigiert',3,corrected)):
        top = word['top']-7
        bottom = word['top']+56
        image = Image.open(build/f'vorlage-{page}.png')
        scale = image.width / float(pdf.pages[page-1].width)
        left,right=(65,510) if name=='diagramm2-korrigiert' else (42,541)
        image.crop(tuple(round(v*scale) for v in (left,top,right,bottom))).save(build/(name+'.png'))

def task(text):
    p = para(text)
    p.paragraph_format.keep_with_next = True
    return p

def answer(text):
    p = doc.add_paragraph()
    p.add_run('Lösung: ').bold=True
    p.add_run(text)
    return p

def picture(name, caption=None):
    if caption:
        p=para(caption,True)
        p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(build/(name+'.png')),width=Mm(170))
    pic=p.runs[0]._r.xpath('.//wp:docPr')[0]
    pic.set('descr',caption or 'Korrigierte Skizze mit direkter n:m-Beziehung zwischen Berechtigungen und Apps')

doc.add_paragraph('Musterlösung Informatik Datenbanken', 'Title')
para('Klasse 10____    Datum: 27.10.2026    Name: ____________________')
para('Version 2 · Gesamt 29 BE')

heading('1 Grundwissen')
para('4 BE',True)
task('Vervollständige die Sätze, indem zu jedem Begriff aus dem Kontext Datenbanken den passenden Begriff aus dem Kontext Programmierung findest.')
for stem, result in (
    ('Eine Spalte entspricht in der objektorientierten Programmierung ', 'einem Attribut'),
    ('Ein Datensatz entspricht in der objektorientierten Programmierung ', 'einem Objekt'),
    ('Eine Tabelle entspricht in der objektorientierten Programmierung ', 'einer Klasse'),
    ('Eine Zelle - außer in der ersten Zeile - entspricht in der objektorientierten Programmierung ', 'einem Attributwert'),
):
    p=doc.add_paragraph(stem)
    p.add_run(result).bold=True
    p.add_run('.')

heading('2 Apps auf dem Handy')
para('11 BE',True)
picture('diagramm1','Klassendiagramm Nr. 1')
task('a) Ergänze das Klassendiagramm, indem du Beziehungen (zwischen den Klassen App und Handy bzw. Handy und Benutzer) ergänzt. Wähle sinnvolle Kardinalitäten und benenne jede Beziehungen eindeutig.')
p=para('Hinweis: Ein Handy gehört in diesem Szenario zu genau einer Person. Eine Person kann aber mehrere Handys haben.')
for r in p.runs: r.italic=True
answer('Die ergänzten Beziehungen sind im Klassendiagramm Nr. 1 dargestellt: App–Handy n:m („ist installiert auf“); Handy–Benutzer n:1 („gehört zu“).')
task('b) Benenne die Beziehungen zwischen den Klassen und begründe deine Entscheidung.')
answer('App–Handy: n:m-Beziehung. Eine App kann auf mehreren Handys installiert sein; ein Handy kann mehrere Apps enthalten.')
para('Benutzer–Handy: 1:n-Beziehung. Eine Person kann mehrere Handys haben; jedes Handy gehört genau einer Person. Daher steht im Diagramm n bei Handy und 1 bei Benutzer.')
task('c) Erkläre bei welcher Beziehungsart eine Beziehungstabelle notwendig ist. Begründe deine Entscheidung.')
answer('Eine Beziehungstabelle ist bei einer n:m-Beziehung notwendig. Ein einzelner Fremdschlüssel in einer der beiden Tabellen reicht für die vielen Zuordnungen in beiden Richtungen nicht aus. Jede Zeile der Beziehungstabelle verbindet über zwei Fremdschlüssel genau zwei Datensätze.')

doc.add_page_break()
heading('3 Berechtigungen und Apps')
para('8 BE',True)
para('Betrachte die Ausschnitte aus den Tabellen Berechtigung, Apps und Berechtigung_von_Apps.')
para('Tabellenausschnitte Nr. 1',True)
for name, heads, rows, widths in (
    ('Berechtigung',['id','name'],[(1,'Kamera'),(2,'Standort'),(3,'Fotos lesen')],[20,70]),
    ('Apps',['id','name'],[(1,'WhatsApp'),(2,'TikTok'),(3,'Instagram')],[20,70]),
    ('Berechtigungen_von_Apps',['app_id','berechtigung_id'],[(1,1),(1,3),(2,1),(3,2),(3,3)],[45,45]),
):
    p=para(name,True)
    p.paragraph_format.keep_with_next=True
    table(heads,rows,widths)
task('a) WhatsApp fragt beim Öffnen der App nach einer neuen Berechtigung "Standort". Du nimmst die Berechtigungsanfrage an.')
p=doc.add_paragraph('Beschreibe die Veränderungen, die passieren müssen, damit die Tabellen (siehe Tabellenausschnitte Nr. 1) auf dem aktuellen Stand sind, ')
p.add_run('genau').bold=True
p.add_run('.')
p.paragraph_format.keep_with_next=True
answer('WhatsApp hat die id 1; Standort hat die id 2. In Berechtigungen_von_Apps wird eine neue Zeile mit app_id = 1 und berechtigung_id = 2 ergänzt. Beide Stammdatensätze existieren bereits; in Apps und Berechtigung sind keine neuen Zeilen erforderlich.')

task('Betrachte das folgende Klassendiagramm mit den Klassen Berechtigungen, Berechtigungen-von-Apps und Apps (siehe Klassendiagramm Nr. 2). Es sind Fehler in diesem Klassendiagramm enthalten.')
picture('diagramm2','Klassendiagramm Nr. 2')
task('b) Skizziere die korrigierte Version des abgebildeten Klassendiagramm, sodass es den Sachverhalt korrekt darstellt.')
picture('diagramm2-korrigiert','Lösung')

task('c) Korrigiere die folgende SQL-Anweisung, indem du falsche Zeilen durchstreichst und in der zweiten Spalte korrigiert angibst.')
t=table(['Vorgegebene SQL-Anweisung','Korrektur'],[
    ('SELECT *','SELECT *'),
    ('FROM Berechtigung, Apps,\nBerechtigung_von_Apps','FROM Berechtigung, Apps,\nBerechtigungen_von_Apps'),
    ('WHERE Berechtigungen.id = Apps.id\nAND Berechtigung_von_Apps.id = App.id','WHERE Berechtigung.id =\nBerechtigungen_von_Apps.berechtigung_id\nAND Apps.id =\nBerechtigungen_von_Apps.app_id;'),
],[85,85])
for row_i,row in enumerate(t.rows[1:],1):
    for cell_i,cell in enumerate(row.cells):
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name='Consolas'
                r.font.size=Pt(8.5)
                if cell_i==0 and row_i>=2: r.font.strike=True

heading('4 Appentwickler')
para('6 BE',True)
para('Betrachte die folgenden Tabellenschemata der Tabellen Appentwickler und App.')
code('Appentwickler (id: int [PS], name: varchar(255),\n    gruendungsjahr: int)')
code('App (id: int [PS], appentwickler_id: int [FS],\n    name: varchar(255), erscheinungsjahr: int,\n    kategorie: varchar(255))')
para('Verwendete Abkürzungen: PS = Primärschlüssel, FS = Fremdschlüssel. App.appentwickler_id verweist auf Appentwickler.id.')
task('a) Beschreibe folgende SQL-Anweisung in eigenen Worten.')
code('SELECT COUNT(*)\nFROM App\nWHERE erscheinungsjahr > 2025')
answer('Die Anweisung zählt die Apps, die nach dem Jahr 2025 erschienen sind, also ab 2026. Ausgegeben wird eine einzelne Zahl. Apps mit Erscheinungsjahr 2025 werden nicht mitgezählt.')
task('b) Formuliere eine SQL-Anweisung, die Namen der Apps zusammen mit den Namen der Appentwickler ausgibt, die die jeweilige App entwickelt haben.')
para('Lösung:',True)
code('SELECT App.name, Appentwickler.name\nFROM App, Appentwickler\nWHERE App.appentwickler_id = Appentwickler.id;')
para('Die beiden Tabellen werden über den Fremdschlüssel App.appentwickler_id und den Primärschlüssel Appentwickler.id verknüpft.')
para('Gesamt 29 BE',True)
para('Viel Erfolg! Franke')

out=root/'Inf10-KA-Version2-loesungen-vollstaendig.docx'
doc.save(out)
from zipfile import ZipFile
from lxml import etree
with ZipFile(out) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','v':'urn:schemas-microsoft-com:vml'}
    assert not xml.findall('.//w:txbxContent',ns)
    assert not xml.findall('.//v:textbox',ns)
    assert len([n for n in z.namelist() if n.startswith('word/media/')])==3
    text=' '.join(xml.xpath('.//w:t/text()',namespaces=ns))
    for phrase in ('Vervollständige die Sätze','Ein Datensatz entspricht','Wähle sinnvolle Kardinalitäten','auf dem aktuellen Stand sind,','Es sind Fehler in diesem Klassendiagramm enthalten.','Verwendete Abkürzungen','Gesamt 29 BE'):
        assert phrase in text,phrase
    assert len(doc.tables)==4
print(out.resolve())
